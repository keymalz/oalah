#!/usr/bin/env python3
"""Membangun laporan .docx dari file nyata (kode, hasil tes JSON/TXT, screenshot).
Jalankan: python3 docs/build_report.py  ->  docs/Laporan-Login-Register-Node-Flutter-Gemini.docx
"""
import json
import os
import re
import struct

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
SHOTS = os.path.join(ROOT, 'screenshots')
RES = os.path.join(ROOT, 'docs', 'test-results')
OUT = os.path.join(ROOT, 'docs', 'Laporan-Login-Register-Node-Flutter-Gemini.docx')
TITLE = 'Implementasi Login dan Register pada Node.js dan Flutter Web dengan Google Sign-In serta Gemini API'
TANGGAL = '14 Oktober 2026'

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for side in ('left_margin', 'right_margin'):
    setattr(sec, side, Cm(2.3))
sec.top_margin = sec.bottom_margin = Cm(2.2)

st = doc.styles['Normal']
st.font.name = 'Arial'
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.25
for name, size in (('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11.5)):
    h = doc.styles[name]
    h.font.name = 'Arial'
    h.font.size = Pt(size)
    h.font.bold = True
    h.font.color.rgb = RGBColor(0x1F, 0x2A, 0x6B)
    h.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    h.paragraph_format.space_before = Pt(14 if name == 'Heading 1' else 10)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True


# ---------- helper ----------
def shade(cell_or_par, hex_fill):
    el = cell_or_par._tc if hasattr(cell_or_par, '_tc') else cell_or_par._p
    pr = el.get_or_add_tcPr() if hasattr(cell_or_par, '_tc') else el.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    pr.append(shd)


def para(text='', bold=False, italic=False, size=None, align=None, color=None, after=None, style=None):
    p = doc.add_paragraph(style=style)
    if text:
        add_runs(p, text, bold, italic, size, color)
    if align:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def add_runs(p, text, bold=False, italic=False, size=None, color=None):
    # **tebal** dan `kode` inline
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not part:
            continue
        b, mono = bold, False
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            part, b = part[2:-2], True
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            part, mono = part[1:-1], True
        r = p.add_run(part)
        r.bold = b
        r.italic = italic
        if mono:
            r.font.name = 'Courier New'
            r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Courier New')
        if size:
            r.font.size = Pt(size - 0.5 if mono else size)
        elif mono:
            r.font.size = Pt(9.5)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)


def bullets(items, style='List Bullet'):
    for n, it in enumerate(items, 1):
        if style == 'List Number':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            add_runs(p, f'{n}.  ' + it)
        else:
            p = doc.add_paragraph(style=style)
            add_runs(p, it)
        p.paragraph_format.space_after = Pt(2)


def code(text):
    for line in text.rstrip('\n').split('\n'):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Cm(0.2)
        r = p.add_run(line if line else ' ')
        r.font.name = 'Courier New'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Courier New')
        r.font.size = Pt(8)
        shade(p, 'F1F3F8')
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def set_borders(table):
    tbl = table._tbl
    pr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for b in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{b}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), '9AA3BF')
        borders.append(el)
    pr.append(borders)


STATUS_FILL = {'Lulus': 'DDF1E4', 'Gagal': 'F9D6D3', 'Belum diuji': 'FFF0C2', 'Selesai': 'DDF1E4', 'Tertunda': 'FFF0C2'}


def fix_layout(t, widths):
    pr = t._tbl.tblPr
    lay = OxmlElement('w:tblLayout')
    lay.set(qn('w:type'), 'fixed')
    pr.append(lay)
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
        for r in t.rows:
            r.cells[i].width = Cm(w)


def table(headers, rows, widths, font=8.5, status_col=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_borders(t)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ''
        add_runs(c.paragraphs[0], h, bold=True, size=font, color='FFFFFF')
        shade(c, '2F3E8F')
    # header berulang
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement('w:tblHeader')
    th.set(qn('w:val'), 'true')
    trPr.append(th)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            add_runs(p, str(v), size=font)
            if status_col is not None and i == status_col and v in STATUS_FILL:
                shade(cells[i], STATUS_FILL[v])
    fix_layout(t, widths)
    for r in t.rows:
        r._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


fig_no = [0]


def figure(fname, caption, width=12.6):
    path = os.path.join(SHOTS, fname)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Cm(width))
    fig_no[0] += 1
    c = para(f'Gambar {fig_no[0]}. {caption}', italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    return c


def missing_box(title, steps):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_borders(t)
    fix_layout(t, [16.4])
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    c = t.rows[0].cells[0]
    shade(c, 'FFF6D6')
    c.text = ''
    add_runs(c.paragraphs[0], f'**BUKTI BELUM TERSEDIA — {title}**', size=9.5)
    for s in steps:
        p = c.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        add_runs(p, '• ' + s, size=9)
    p = c.add_paragraph()
    add_runs(p, '[Tempel screenshot asli di sini setelah langkah dilakukan]', italic=True, size=9, color='7A5B00')
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def read(path):
    with open(path, encoding='utf-8') as f:
        return f.read()


def snippet(rel, start, end=None, include_end=False, max_lines=None):
    src = read(os.path.join(ROOT, rel))
    a = src.index(start)
    b = src.index(end, a + len(start)) + (len(end) if include_end else 0) if end else len(src)
    s = src[a:b].rstrip() + '\n'
    lines = s.split('\n')
    if max_lines and len(lines) > max_lines:
        s = '\n'.join(lines[:max_lines]) + '\n// ... (dipotong)\n'
    return s


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------- data hasil ----------
node_rows = json.load(open(os.path.join(RES, 'node-e2e.json'), encoding='utf-8'))
fl_rows = json.load(open(os.path.join(RES, 'flutter-e2e.json'), encoding='utf-8'))
backend_txt = read(os.path.join(RES, 'node-backend-test.txt'))
backend_ok = re.findall(r'^(ok|not ok) \d+ - (.*)$', backend_txt, re.M)
flutter_txt = read(os.path.join(RES, 'flutter-test.txt'))
widget_names = re.findall(r'\+(\d+): (.*)', flutter_txt)
widget_all_pass = 'All tests passed' in flutter_txt
versions = read(os.path.join(RES, 'versions.txt'))

shots = sorted(os.listdir(SHOTS))
by_prefix = {f.split('-')[0]: f for f in shots}


def S(n):  # nama file berdasarkan nomor
    return by_prefix[f'{n:02d}']


# ---------- SAMPUL ----------
for _ in range(3):
    para('')
para('LAPORAN PRAKTIKUM', bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, color='5B6380', after=10)
para(TITLE, bold=True, size=20, align=WD_ALIGN_PARAGRAPH.CENTER, color='1F2A6B', after=24)
para('Node.js (Express) · Flutter Web · Firebase Authentication · Google Sign-In · Google AI Studio / Gemini API',
     size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, color='5B6380', after=36)
t = doc.add_table(rows=0, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
set_borders(t)
for k in ('Nama', 'NIM', 'Kelas', 'Mata kuliah', 'Dosen pengampu', 'Institusi'):
    cells = t.add_row().cells
    cells[0].text = ''
    add_runs(cells[0].paragraphs[0], k, bold=True)
    shade(cells[0], 'EEF0F8')
    cells[1].text = ''
    add_runs(cells[1].paragraphs[0], '………………………………………………  (isi sendiri)', color='7A7F99')
fix_layout(t, [4.5, 10.5])
para('')
para(f'Tanggal: {TANGGAL}', align=WD_ALIGN_PARAGRAPH.CENTER, after=24)
box = doc.add_table(rows=1, cols=1)
box.alignment = WD_TABLE_ALIGNMENT.CENTER
set_borders(box)
fix_layout(box, [16.4])
c = box.rows[0].cells[0]
shade(c, 'FFF6D6')
c.text = ''
add_runs(c.paragraphs[0],
         '**Catatan status:** seluruh kode dan pengujian otomatis telah selesai dan lulus terhadap **Firebase Auth Emulator** '
         'dan **Gemini mock**. Pembuatan proyek Firebase nyata, API key Google AI Studio nyata, serta login Google nyata '
         '**belum dilakukan** karena memerlukan akun Google pemilik tugas. Lihat Bab 2 dan Bab 13. Dokumen ini adalah berkas .docx '
         'yang siap diunggah/dikonversi ke Google Docs.', size=9.5)
page_break()

# ---------- DAFTAR ISI ----------
doc.add_heading('Daftar Isi', 1)
for i, s in enumerate([
    'Tujuan praktikum', 'Ringkasan hasil dan status penyelesaian', 'Tools, versi, dan arsitektur',
    'Konfigurasi Firebase dan Google Sign-In', 'Pembuatan API key melalui Google AI Studio',
    'Implementasi website Node.js', 'Implementasi website Flutter Web', 'Integrasi Gemini dan pengamanan key',
    'Langkah menjalankan kedua aplikasi', 'Demonstrasi fitur dengan screenshot dan caption', 'Tabel pengujian',
    'Kendala nyata dan penyelesaiannya', 'Kesimpulan dan keterbatasan', 'Referensi dokumentasi resmi'], start=2):
    para(f'{i - 1}. {s}', after=1)
page_break()

# ---------- 1. TUJUAN ----------
doc.add_heading('1. Tujuan Praktikum', 1)
bullets([
    'Membangun website **login/register** dengan **Node.js** (Express) yang mendukung email/password dan Google, '
    'dashboard terproteksi, dan logout.',
    'Membangun website **login/register** serupa dengan **Flutter Web**.',
    'Memakai **Firebase Authentication** untuk pengelolaan kredensial, dan **Firebase Admin SDK** di backend untuk memverifikasi identitas.',
    'Membuat **API key Gemini** melalui Google AI Studio dan mengintegrasikannya secara aman (key hanya di server) '
    'melalui fitur **"Tanya Gemini"** pada dashboard kedua website.',
    'Mendokumentasikan implementasi, pengujian, dan bukti (screenshot) dalam satu dokumen.',
])

# ---------- 2. RINGKASAN ----------
doc.add_heading('2. Ringkasan Hasil dan Status Penyelesaian', 1)
para('Tabel berikut membedakan hasil **aktual**, pengujian dengan **emulator/mock**, dan bagian yang **belum dapat diverifikasi**.')
table(['Komponen', 'Status', 'Keterangan'], [
    ['Website Node.js + backend (register, login, dashboard, logout, Tanya Gemini)', 'Selesai',
     'Kode selesai; 11/11 tes backend dan 17 skenario E2E lulus (emulator + mock).'],
    ['Website Flutter Web (fitur sama)', 'Selesai',
     '`flutter analyze` bersih; 5/5 tes unit+widget dan 14 skenario E2E lulus (emulator + mock).'],
    ['Endpoint `POST /api/gemini` (validasi, autentikasi, penanganan error)', 'Selesai',
     'Teruji, termasuk panggilan ke Gemini API nyata dengan key palsu (dipetakan ke `GEMINI_BAD_KEY`).'],
    ['Proyek Firebase nyata + provider Email/Password & Google aktif', 'Tertunda',
     'Perlu login Google pemilik. Konsol Firebase tidak dapat diakses dari lingkungan pengembangan.'],
    ['API key Google AI Studio nyata + jawaban Gemini nyata (Node & Flutter)', 'Tertunda',
     'Perlu login Google pemilik. Jawaban pada screenshot adalah **MOCK**, bukan Gemini asli.'],
    ['Login Google **nyata** (popup OAuth) pada kedua website', 'Tertunda',
     'Belum diuji. Pengujian memakai kredensial IdP tiruan milik emulator (bukan Google asli).'],
    ['Google Docs asli', 'Tertunda', 'Dokumen ini berformat .docx; lihat catatan konversi pada bagian akhir bab ini.'],
], [6.2, 2.2, 7.6], status_col=1)
para('**Bukti yang masih kurang** (harus diambil setelah konfigurasi nyata): (1) layar Firebase yang memperlihatkan provider '
     'Email/Password dan Google aktif; (2) daftar pengguna uji di Firebase Console; (3) pembuatan API key di AI Studio '
     '(nilai key disamarkan); (4) jawaban Gemini nyata di Node.js; (5) jawaban Gemini nyata di Flutter Web; '
     '(6) dashboard setelah login Google nyata di kedua website; (7) pembatalan popup Google nyata.')
para('**Catatan konversi ke Google Docs:** unggah berkas .docx ke Google Drive → klik kanan → *Buka dengan → Google Dokumen*, '
     'atau *File → Simpan sebagai Google Dokumen*. Judul dan screenshot tertanam akan terbawa.')

# ---------- 3. TOOLS ----------
doc.add_heading('3. Tools, Versi, dan Arsitektur', 1)
doc.add_heading('3.1 Tools dan versi yang digunakan', 2)
para('Versi berikut dicatat dari lingkungan yang benar-benar dipakai (`docs/test-results/versions.txt`).')
table(['Komponen', 'Versi / keterangan'], [
    ['Node.js / npm', 'v22.22.0 / 10.9.4 (LTS; firebase-admin 14 mensyaratkan Node ≥ 22)'],
    ['Express', '5.3.0 (+ helmet 8.3.0, cookie-parser 1.4.7, express-rate-limit 8.7.1, ejs 7.0.1)'],
    ['Firebase Admin SDK (backend)', 'firebase-admin 14.5.0 (API modular `firebase-admin/app`, `firebase-admin/auth`)'],
    ['Firebase JS SDK (browser Node.js)', 'firebase 13.0.0, dibundel dengan esbuild 0.28.2'],
    ['SDK Gemini', '@google/genai 2.29.1 (SDK JavaScript resmi Google Gen AI)'],
    ['Flutter / Dart', 'Flutter 3.47.7 (stable) / Dart 3.13.5'],
    ['Paket Flutter', 'firebase_core ^4.15.0, firebase_auth ^6.7.0 (firebase_auth_web 6.3.0), http ^1.6.0'],
    ['Pengujian', 'node:test + supertest 7.3.1; flutter_test; Playwright 1.56.1 (Chromium); firebase-tools 15.33.0 (Auth Emulator), OpenJDK 21'],
    ['Model Gemini', 'Dikonfigurasi lewat `GEMINI_MODEL` (default `gemini-3.5-flash`). **Belum diverifikasi terhadap akun** — '
                     'pilih model yang tampil di AI Studio akun Anda.'],
], [5, 11])
doc.add_heading('3.2 Arsitektur', 2)
para('**Node.js menjalankan server dan backend, sedangkan Firebase Authentication mengelola kredensial pengguna.** '
     'Aplikasi ini tidak menyimpan password; password dikelola dan di-hash oleh Firebase.')
code('''Browser (Node.js)                       Browser (Flutter Web)
  Firebase JS SDK  -- login -->  Firebase Auth  <-- login --  firebase_auth
        | ID token                                              | ID token
        v                                                       v
  POST /api/sessionLogin (+CSRF)                Authorization: Bearer <ID token>
        | Admin SDK: verifyIdToken                              |
        v                                                       |
  Set-Cookie __session (HttpOnly, SameSite=Lax) -> /dashboard   |
        |                                                       |
        +------------> Node.js / Express backend <--------------+
                       verifikasi cookie ATAU Bearer (Admin SDK)
                       POST /api/gemini -> Gemini API (GEMINI_API_KEY hanya di server)''')
bullets([
    '**Website Node.js:** SDK Firebase di browser menangani register/login. Backend menukar ID token yang valid (dan baru, < 5 menit) '
    'menjadi **session cookie** `HttpOnly`, `SameSite=Lax`, `Secure` pada production/HTTPS. `/dashboard` dan `/api/*` memverifikasi cookie. '
    'Permintaan POST berbasis cookie dilindungi **token CSRF** dan pemeriksaan header `Origin`.',
    '**Flutter Web:** `firebase_auth` menangani login; ID token dikirim sebagai **Bearer** ke backend, yang memverifikasinya sebelum memanggil Gemini.',
    '**CORS** dibatasi ke origin pada `ALLOWED_ORIGINS`; alamat backend Flutter dapat dikonfigurasi lewat `--dart-define=API_BASE_URL`.',
    '**Satu proyek Firebase** direncanakan untuk kedua aplikasi.',
])

# ---------- 4. FIREBASE ----------
doc.add_heading('4. Konfigurasi Firebase dan Google Sign-In', 1)
para('**Status: belum dilakukan pada Firebase nyata.** Lingkungan pengembangan tidak memiliki login Google dan memblokir '
     '`console.firebase.google.com` serta `apis.google.com`. Langkah di bawah adalah prosedur yang harus dijalankan pemilik akun; '
     'seluruh kode sudah menerima konfigurasinya lewat variabel lingkungan/`--dart-define`.')
doc.add_heading('4.1 Langkah konfigurasi nyata', 2)
bullets([
    'Buka Firebase Console → **Add project** (paket Spark/gratis, jangan aktifkan billing).',
    '**Authentication → Get started → Sign-in method:** aktifkan **Email/Password** dan **Google** (isi *support email*).',
    '**Authentication → Settings → Authorized domains:** pastikan `localhost` tercantum dan tambahkan domain deploy (jika ada).',
    '**Project settings → Your apps → Web (`</>`):** daftarkan satu aplikasi web, salin `apiKey`, `authDomain`, `projectId`, `appId`, '
    '`messagingSenderId` (nilai publik) ke `node-web/.env` dan `flutter-web/dart-defines.json`.',
    '**Project settings → Service accounts → Generate new private key** → simpan ke `node-web/secrets/serviceAccount.json` '
    '(sudah ada di `.gitignore`). Jangan pernah mengirim file ini lewat chat/screenshot.',
    'Flutter: konfigurasi dapat memakai `flutterfire configure` (alur resmi) atau `--dart-define-from-file` seperti pada proyek ini.',
], style='List Number')
doc.add_heading('4.2 Bukti yang diperlukan', 2)
missing_box('Provider autentikasi aktif (Firebase Console)', [
    'Buka Authentication → Sign-in method; tangkap layar daftar provider yang menampilkan Email/Password dan Google berstatus Enabled.'])
missing_box('Daftar pengguna uji (Firebase Console)', [
    'Buka Authentication → Users setelah menguji kedua website; samarkan email pribadi (gunakan akun uji).'])
para('Sebagai pembanding **bukan pengganti**, daftar pengguna pada **Firebase Auth Emulator UI** (lokal, data uji) ditampilkan pada Gambar di Bab 10.', italic=True)

# ---------- 5. AI STUDIO ----------
doc.add_heading('5. Pembuatan API Key melalui Google AI Studio', 1)
para('**Status: belum dilakukan.** `aistudio.google.com` tidak dapat dijangkau dari lingkungan pengembangan dan pembuatan key memerlukan login Google. '
     'Integrasi server sudah siap dan hanya membutuhkan nilai `GEMINI_API_KEY` dan `GEMINI_MODEL` pada `.env`.')
bullets([
    'Buka <https://aistudio.google.com/apikey> → **Create API key** (pilih/buat proyek Google Cloud; gunakan tier gratis, jangan aktifkan billing).',
    'Salin key ke `node-web/.env` → `GEMINI_API_KEY=...`. **Jangan** menaruhnya di kode browser, Flutter, `--dart-define`, Git, atau screenshot.',
    'Pada AI Studio, lihat model yang tersedia untuk akun Anda dan isi `GEMINI_MODEL`.',
    'Untuk screenshot laporan, **samarkan seluruh nilai key** (tampilkan hanya nama key dan 4 karakter terakhir bila perlu).',
], style='List Number')
missing_box('Pembuatan API key Google AI Studio (nilai key disamarkan)', [
    'Tangkap layar halaman API keys setelah key dibuat; tutup/blur seluruh nilai key.'])

# ---------- 6. NODE ----------
doc.add_heading('6. Implementasi Website Node.js', 1)
para('Struktur: `node-web/server.js`, `src/` (config, app, auth, csrf, gemini, firebaseAdmin), `views/*.ejs`, `client/*.js` '
     '(dibundel esbuild menjadi `public/js/*.bundle.js`), `test/*.test.js`.')
doc.add_heading('6.1 Alur autentikasi', 2)
bullets([
    'Halaman `/register` dan `/login` memvalidasi input di sisi klien (email, panjang password ≥ 8 dengan huruf dan angka, konfirmasi cocok), '
    'mencegah submit ganda, dan menampilkan pesan error berbahasa Indonesia untuk kode error Firebase.',
    'SDK Firebase (persistensi in-memory) melakukan `createUserWithEmailAndPassword`/`signInWithEmailAndPassword`/`signInWithPopup(GoogleAuthProvider)`.',
    'ID token dikirim ke `POST /api/sessionLogin` bersama header `X-CSRF-Token`; server memverifikasi token lalu membuat session cookie.',
    'Setelah itu SDK klien `signOut()` — sesi dipegang server lewat cookie `HttpOnly`.',
    '`GET /dashboard` memverifikasi cookie (`verifySessionCookie(..., true)`); gagal → hapus cookie dan arahkan ke `/login` '
    '(`?expired=1` bila sesi kedaluwarsa).',
])
doc.add_heading('6.2 Cuplikan kode inti', 2)
para('Penukaran ID token menjadi session cookie (`node-web/src/app.js`):', bold=True, after=2)
code(snippet('node-web/src/app.js', "  app.post('/api/sessionLogin'", "  app.post('/api/sessionLogout'"))
para('Verifikasi identitas — Bearer (Flutter) atau cookie (Node) (`node-web/src/auth.js`):', bold=True, after=2)
code(snippet('node-web/src/auth.js', '// Mengembalikan { user, via }', 'function describeUser'))
para('Perlindungan CSRF (`node-web/src/csrf.js`):', bold=True, after=2)
code(snippet('node-web/src/csrf.js', '// Pola "synchronizer token"', 'function safeEqual'))
para('Penukaran sesi di browser (`node-web/client/session.js`):', bold=True, after=2)
code(snippet('node-web/client/session.js', '// Menukar user Firebase'))
doc.add_heading('6.3 Pengamanan', 2)
bullets([
    'Cookie sesi: `HttpOnly`, `SameSite=Lax`, `Secure` saat `NODE_ENV=production`; masa berlaku dapat diatur (`SESSION_MAX_DAYS`).',
    'CSRF: token acak 256 bit di cookie `HttpOnly` + meta tag halaman; dibandingkan dengan header `X-CSRF-Token` (timing-safe). Paket `csurf` tidak dipakai karena sudah diarsipkan.',
    'Header keamanan via `helmet` (CSP ketat: skrip hanya dari `self` dan `apis.google.com`; `frame-ancestors none`).',
    'Rate limit `POST /api/gemini`: 10 permintaan/menit per pengguna. Body JSON dibatasi 10 KB.',
])

# ---------- 7. FLUTTER ----------
doc.add_heading('7. Implementasi Website Flutter Web', 1)
para('Struktur: `flutter-web/lib/` — `main.dart` (inisialisasi Firebase), `app.dart` (AuthGate), `config.dart`, `firebase_options.dart`, '
     '`services/auth_repository.dart`, `services/gemini_client.dart`, `pages/{login,register,dashboard}_page.dart`, `validators.dart`, `widgets.dart`.')
bullets([
    '`AuthRepository` membungkus Firebase Auth sehingga UI dapat diuji dengan implementasi palsu (dipakai pada tes widget).',
    '`AuthGate` mendengarkan `userChanges()`: tanpa pengguna menampilkan login/register; dengan pengguna menampilkan dashboard. '
    'Semua URL dialihkan melalui `AuthGate`, sehingga `/#/dashboard` tidak membuka dashboard tanpa login.',
    'Login Google memakai `FirebaseAuth.signInWithPopup(GoogleAuthProvider())`; pembatalan popup (`popup-closed-by-user`, `cancelled-popup-request`) ditangani dengan pesan khusus.',
    'Konfigurasi publik Firebase dan alamat backend diberikan lewat `--dart-define`; **tidak ada API key Gemini di Flutter**.',
    'Persistensi sesi bawaan web (IndexedDB) membuat sesi bertahan saat halaman di-refresh.',
])
para('Cuplikan kode inti (`auth_repository.dart`) — register dan login Google:', bold=True, after=2)
code(snippet('flutter-web/lib/services/auth_repository.dart', '  Future<T> _guard<T>', '  @override\n  Future<void> signOut'))
para('Pemanggilan backend dengan Bearer token (`gemini_client.dart`):', bold=True, after=2)
code(snippet('flutter-web/lib/services/gemini_client.dart', '  Future<GeminiAnswer> ask', '    Map<String, dynamic> body'))

# ---------- 8. GEMINI ----------
doc.add_heading('8. Integrasi Gemini dan Pengamanan Key', 1)
para('Endpoint `POST /api/gemini` menerima `{prompt}`, memeriksa autentikasi (cookie+CSRF atau Bearer), memvalidasi input '
     '(tidak kosong, ≤ 2000 karakter), membatasi laju, lalu memanggil Gemini lewat `@google/genai`. Konfigurasi: `GEMINI_API_KEY`, '
     '`GEMINI_MODEL`, `GEMINI_TIMEOUT_MS`, `GEMINI_MAX_INPUT_CHARS`.')
code(snippet('node-web/src/gemini.js', 'async function askGemini', 'module.exports'))
doc.add_heading('8.1 Penanganan kegagalan', 2)
table(['Kondisi', 'Kode/HTTP ke klien', 'Pesan'], [
    ['Tidak login / token tidak valid', '401 `NO_AUTH`/`TOKEN_INVALID`/`SESSION_EXPIRED`', 'Anda belum masuk atau token tidak valid.'],
    ['Cookie tanpa token CSRF', '403 `CSRF_INVALID`', 'Token keamanan tidak valid.'],
    ['Prompt kosong / terlalu panjang', '400 `EMPTY_PROMPT` / `PROMPT_TOO_LONG`', 'Pertanyaan tidak boleh kosong / maksimal N karakter.'],
    ['Key tidak ada / salah / ditolak', '503 `GEMINI_NO_KEY` / `GEMINI_BAD_KEY`', 'Konfigurasi server (tanpa menyebut nilai key).'],
    ['Kuota habis / terlalu banyak permintaan', '429 `GEMINI_QUOTA`', 'Kuota Gemini API habis atau terlalu banyak permintaan.'],
    ['Timeout', '504 `GEMINI_TIMEOUT`', 'Gemini tidak merespons tepat waktu.'],
    ['Kegagalan lain / jawaban kosong', '502 `GEMINI_FAILED`/`GEMINI_EMPTY`', 'Gemini gagal memproses permintaan.'],
], [5, 5.4, 5.6])
doc.add_heading('8.2 Pengamanan key', 2)
bullets([
    '`GEMINI_API_KEY` hanya dibaca server dari `.env` (ada di `.gitignore`); `.env.example` berisi placeholder.',
    'Key tidak pernah dikirim ke browser maupun dimasukkan ke build Flutter (`--dart-define` hanya berisi konfigurasi publik Firebase dan alamat backend).',
    'Detail error asli dari Google tidak diteruskan ke klien dan log hanya mencatat kode + UID, bukan prompt atau key. '
    'Tes otomatis memastikan string rahasia tidak muncul pada respons error.',
])

# ---------- 9. MENJALANKAN ----------
doc.add_heading('9. Langkah Menjalankan Kedua Aplikasi', 1)
doc.add_heading('9.1 Konfigurasi nyata (setelah Bab 4 dan 5)', 2)
code('''# Node.js + backend
cd node-web
npm install
cp .env.example .env        # isi FIREBASE_*, GOOGLE_APPLICATION_CREDENTIALS, GEMINI_API_KEY, GEMINI_MODEL
npm run build:client
npm start                   # http://localhost:3000

# Flutter Web (terminal lain)
cd flutter-web
cp dart-defines.example.json dart-defines.json   # isi konfigurasi web Firebase
flutter pub get
flutter run -d chrome --web-port 8080 --dart-define-from-file=dart-defines.json''')
doc.add_heading('9.2 Perintah yang benar-benar dipakai pada pengujian lokal (emulator + mock)', 2)
code('''firebase emulators:start --only auth --project demo-auth-gemini

cd node-web && FIREBASE_AUTH_EMULATOR_HOST=127.0.0.1:9099 npm test
FIREBASE_AUTH_EMULATOR_HOST=127.0.0.1:9099 GEMINI_MOCK=1 ALLOWED_ORIGINS=http://localhost:8080 node server.js

cd flutter-web && flutter analyze && flutter test
flutter build web --release --no-web-resources-cdn \\
  --dart-define=USE_AUTH_EMULATOR=true --dart-define=API_BASE_URL=http://localhost:3000
(cd build/web && python3 -m http.server 8080 --bind 127.0.0.1)

cd docs/e2e && npm install && npm run build-shim
node e2e-node.js && SHOT_START=12 node e2e-flutter.js && SHOT_START=23 node console-shots.js''')
doc.add_heading('9.3 Urutan demonstrasi', 2)
bullets([
    'Register (nama, email, password, konfirmasi) di Node.js → dashboard → **Tanya Gemini** → **Keluar**.',
    'Buka `/dashboard` langsung → dialihkan ke `/login` (terproteksi).',
    'Login email/password → dashboard; keluar; **Masuk dengan Google** → dashboard (metode login: Google).',
    'Ulangi seluruh urutan di Flutter Web (`http://localhost:8080`).',
], style='List Number')

# ---------- 10. DEMONSTRASI ----------
doc.add_heading('10. Demonstrasi Fitur dengan Screenshot dan Caption', 1)
para('Seluruh gambar diambil dari aplikasi yang benar-benar dijalankan oleh Playwright/Chromium. **Label sumber** pada tiap caption: '
     '**[Emulator]** = Firebase Auth Emulator lokal (bukan Firebase Console/akun nyata); **[Mock]** = jawaban Gemini tiruan; '
     '**[Konsol]** = output perintah nyata yang dirender menjadi gambar. Data uji bersifat fiktif (`@example.com`).')

CAP = {
    1: ('10.1 Website Node.js', 'Halaman register Node.js (`/register`): nama, email, password, konfirmasi, tombol Daftar, tombol Google, tautan ke login. [Emulator]'),
    2: (None, 'Validasi register Node.js: nama terlalu pendek, email tidak valid, password lemah; kolom bermasalah ditandai dan fokus dipindah ke kolom pertama. [Emulator]'),
    3: (None, 'Password dan konfirmasi tidak cocok pada register Node.js: muncul "Konfirmasi password tidak cocok." [Emulator]'),
    4: (None, 'Dashboard Node.js setelah register email/password: nama, email, metode login "Email dan password", form Tanya Gemini. [Emulator]'),
    5: (None, 'Tanya Gemini pada Node.js. **Jawaban ini MOCK (berlabel), BUKAN bukti Gemini nyata.** [Mock]'),
    6: (None, 'Setelah logout, membuka `/dashboard` langsung dialihkan ke halaman login Node.js. [Emulator]'),
    7: (None, 'Login Node.js dengan password salah: "Email atau password salah." [Emulator]'),
    8: (None, 'Dashboard Node.js setelah login email/password yang benar. [Emulator]'),
    9: (None, 'Sesi hilang saat memakai Gemini: API membalas 401 dan pengguna dialihkan ke login dengan pesan sesi berakhir. [Emulator]'),
    10: (None, 'Register Node.js dengan email yang sudah dipakai: "Email ini sudah terdaftar. Silakan masuk." [Emulator]'),
    11: (None, 'Dashboard Node.js untuk pengguna dengan provider **google.com**. **Ini kredensial IdP tiruan emulator, BUKAN login Google nyata.** [Emulator]'),
    12: ('10.2 Website Flutter Web', 'Halaman login Flutter Web. Pita merah di bawah adalah penanda bawaan Firebase bahwa Auth Emulator aktif. [Emulator]'),
    13: (None, 'Mengakses `/#/dashboard` tanpa login di Flutter Web: yang tampil tetap halaman login (AuthGate). [Emulator]'),
    14: (None, 'Validasi login Flutter Web: email kosong/tidak valid dan password kosong. [Emulator]'),
    15: (None, 'Halaman register Flutter Web. [Emulator]'),
    16: (None, 'Validasi register Flutter Web: nama pendek, email tidak valid, password lemah, konfirmasi tidak cocok. [Emulator]'),
    17: (None, 'Dashboard Flutter Web setelah register email/password. [Emulator]'),
    18: (None, 'Tanya Gemini pada Flutter Web melalui backend Node.js (Bearer token). **Jawaban MOCK, BUKAN Gemini nyata.** [Mock]'),
    19: (None, 'Login Flutter Web dengan password salah. [Emulator]'),
    20: (None, 'Dashboard Flutter Web setelah login email/password. [Emulator]'),
    21: (None, 'Register Flutter Web dengan email yang sudah dipakai. [Emulator]'),
    22: (None, 'Dashboard Flutter Web untuk provider **google.com**. **Kredensial IdP tiruan emulator, BUKAN login Google nyata.** [Emulator]'),
    23: ('10.3 Bukti konsol dan emulator', 'Hasil `npm test` backend (11 tes): autentikasi, CSRF, sesi, CORS, pemetaan error Gemini, dan panggilan ke Gemini API nyata dengan key palsu. [Konsol]'),
    24: (None, '`flutter analyze`, `flutter test` (5 tes unit+widget), dan `flutter build web`. [Konsol]'),
    25: (None, 'Ringkasan skenario E2E website Node.js. [Konsol]'),
    26: (None, 'Ringkasan skenario E2E Flutter Web. [Konsol]'),
    27: (None, 'Daftar pengguna uji pada **Firebase Auth Emulator UI** (ikon surat = email/password, ikon G = provider Google tiruan). '
               '**Ini bukan Firebase Console.** [Emulator]'),
}
for n in sorted(CAP):
    head, cap = CAP[n]
    if head:
        doc.add_heading(head, 2)
    figure(S(n), cap)

doc.add_heading('10.4 Bukti yang belum dapat diambil', 2)
missing_box('Dashboard setelah login Google NYATA — Node.js', ['Aktifkan provider Google, jalankan Node.js, klik "Masuk dengan Google" dengan akun uji, tangkap dashboard.'])
missing_box('Dashboard setelah login Google NYATA — Flutter Web', ['Sama seperti di atas pada `http://localhost:8080`.'])
missing_box('Jawaban Gemini NYATA — Node.js dan Flutter Web', ['Setelah `GEMINI_API_KEY` valid, ajukan pertanyaan di kedua dashboard; tangkap jawabannya (tanpa key).'])
missing_box('Pembatalan popup Google NYATA', ['Klik tombol Google lalu tutup popup; tangkap pesan "Login Google dibatalkan..." pada kedua website.'])

# ---------- 11. TABEL PENGUJIAN ----------
doc.add_heading('11. Tabel Pengujian', 1)
para('Status: **Lulus** (diuji dan sesuai), **Gagal**, **Belum diuji**. Kolom "Lingkungan" menunjukkan dasar bukti. '
     'Keberhasilan di emulator **tidak** membuktikan Google Sign-In nyata berhasil.')
doc.add_heading('11.1 Tes backend (node:test + supertest, Auth Emulator)', 2)
BACKEND_EXP = {
    'dashboard tanpa login diarahkan ke /login': ('302 ke /login', 'Gambar 23'),
    'POST /api/gemini tanpa autentikasi ditolak 401 dan Gemini tidak dipanggil': ('401, Gemini tidak dipanggil', 'Gambar 23'),
    'bearer token palsu ditolak 401': ('401 TOKEN_INVALID', 'Gambar 23'),
    'session cookie palsu ditolak 401 dan cookie dihapus': ('401 dan cookie dihapus', 'Gambar 23'),
    'sessionLogin tanpa CSRF ditolak 403': ('403 CSRF_INVALID', 'Gambar 23'),
    'sessionLogin dengan ID token palsu ditolak 401': ('401', 'Gambar 23'),
    'alur sesi: login -> dashboard -> gemini -> logout': ('Cookie HttpOnly+SameSite=Lax; dashboard 200; prompt kosong/panjang ditolak 400; jawaban 200; logout menghapus cookie', 'Gambar 23'),
    'bearer token valid (gaya Flutter) diterima': ('200 dan /api/me berisi email + provider password', 'Gambar 23'),
    'error Gemini dipetakan dan tidak membocorkan rahasia': ('Kode 503/429/504/502 sesuai; string rahasia tidak ada di respons', 'Gambar 23'),
    'CORS hanya untuk origin yang diizinkan': ('Origin terdaftar diizinkan; origin lain tanpa header CORS', 'Gambar 23'),
    'key Gemini tidak valid -> 503 GEMINI_BAD_KEY tanpa membocorkan key': ('Panggilan ke Gemini API NYATA dengan key palsu → 503 GEMINI_BAD_KEY, key tidak bocor', 'Gambar 23'),
}
rows = []
for i, (st_, name) in enumerate(backend_ok, 1):
    exp, ev = BACKEND_EXP.get(name, ('-', '-'))
    rows.append([f'B-{i:02d}', name, exp, 'Sesuai' if st_ == 'ok' else 'Menyimpang', 'Lulus' if st_ == 'ok' else 'Gagal', ev])
table(['ID', 'Skenario', 'Hasil diharapkan', 'Hasil aktual', 'Status', 'Bukti'], rows, [1.1, 4.8, 4.6, 1.8, 1.7, 2.0], font=8, status_col=4)

doc.add_heading('11.2 E2E website Node.js (browser, Auth Emulator + Gemini mock)', 2)
EV = {  # id -> gambar
    'N-01': 1, 'N-02': 2, 'N-03': 3, 'N-05': 4, 'N-07': 5, 'N-09': 6, 'N-11': 7, 'N-12': 8, 'N-14': 9, 'N-15': 10, 'N-17': 11,
    'F-01': 12, 'F-02': 13, 'F-03': 14, 'F-04': 15, 'F-05': 16, 'F-06': 17, 'F-07': 18, 'F-10': 19, 'F-11': 20, 'F-12': 21, 'F-13': 22,
}


def e2e_rows(data):
    out = []
    for r in data:
        ev = f'Gambar {EV[r["id"]]}' if r['id'] in EV else 'Log E2E (Gambar 25/26)'
        actual = r['actual'] if r['status'] != 'Belum diuji' else 'Belum diuji: popup Google asli tidak dapat dibuka (apis.google.com diblokir) — perlu uji manual.'
        out.append([r['id'], r['scenario'], r['expected'], actual, r['status'], ev])
    return out


table(['ID', 'Skenario', 'Hasil diharapkan', 'Hasil aktual', 'Status', 'Bukti'], e2e_rows(node_rows), [1.1, 4.2, 3.9, 3.6, 1.6, 1.6], font=8, status_col=4)
doc.add_heading('11.3 E2E Flutter Web (browser, Auth Emulator + backend Node.js + Gemini mock)', 2)
table(['ID', 'Skenario', 'Hasil diharapkan', 'Hasil aktual', 'Status', 'Bukti'], e2e_rows(fl_rows), [1.1, 4.2, 3.9, 3.6, 1.6, 1.6], font=8, status_col=4)
doc.add_heading('11.4 Tes unit dan widget Flutter (repository/HTTP palsu)', 2)
rows = []
for i, nm in enumerate([n for _, n in widget_names if not n.startswith('loading') and not n.startswith('All tests passed')], 1):
    rows.append([f'W-{i:02d}', nm, 'Lulus' if widget_all_pass else 'Gagal', 'Gambar 24'])
table(['ID', 'Skenario', 'Status', 'Bukti'], rows, [1.4, 10.4, 2, 2.2], font=8, status_col=2)
doc.add_heading('11.5 Skenario yang belum diuji (butuh akun/konfigurasi nyata)', 2)
rows_pending = [
    ['R-01', 'Provider Email/Password & Google aktif di Firebase Console', 'Layar provider menunjukkan Enabled', 'Belum dilakukan', 'Belum diuji', '-'],
    ['R-02', 'Daftar pengguna uji di Firebase Console', 'Pengguna uji tampil', 'Hanya Emulator UI (Gambar 27)', 'Belum diuji', '-'],
    ['R-03', 'Pembuatan API key di Google AI Studio', 'Key terbuat (nilai disamarkan)', 'Belum dilakukan', 'Belum diuji', '-'],
    ['R-04', 'Jawaban Gemini NYATA di Node.js', 'Jawaban model asli', 'Belum dilakukan (hanya MOCK)', 'Belum diuji', '-'],
    ['R-05', 'Jawaban Gemini NYATA di Flutter Web', 'Jawaban model asli', 'Belum dilakukan (hanya MOCK)', 'Belum diuji', '-'],
    ['R-06', 'Login Google NYATA (popup OAuth) di Node.js', 'Dashboard, metode Google', 'Belum dilakukan', 'Belum diuji', '-'],
    ['R-07', 'Login Google NYATA (popup OAuth) di Flutter Web', 'Dashboard, metode Google', 'Belum dilakukan', 'Belum diuji', '-'],
    ['R-08', 'Pembatalan popup Google NYATA di kedua website', 'Pesan pembatalan', 'Logika pesan hanya diuji dengan repository palsu (W-04)', 'Belum diuji', '-'],
    ['R-09', 'Sesi benar-benar kedaluwarsa karena waktu (cookie/ID token)', 'Diarahkan ke login dengan pesan', 'Disimulasikan dengan menghapus cookie (N-14) dan 401 palsu (W-05)', 'Belum diuji', '-'],
    ['R-10', 'Timeout dan kuota habis dari Gemini NYATA', '504 / 429', 'Hanya pemetaan error (B-09)', 'Belum diuji', '-'],
    ['R-11', 'Cookie `Secure` pada HTTPS production', 'Flag Secure aktif', 'Belum dideploy (N-13 hanya di HTTP lokal)', 'Belum diuji', '-'],
]
table(['ID', 'Skenario', 'Hasil diharapkan', 'Hasil aktual', 'Status', 'Bukti'], rows_pending, [1.1, 4.4, 3.4, 4.0, 1.6, 1.5], font=8, status_col=4)

# ---------- 12. KENDALA ----------
doc.add_heading('12. Kendala Nyata dan Penyelesaiannya', 1)
table(['No', 'Kendala (benar-benar terjadi)', 'Penyelesaian'], [
    ['1', 'Lingkungan pengembangan memblokir `aistudio.google.com`, `console.firebase.google.com`, `apis.google.com`, dan tidak punya login Google.',
     'Proyek Firebase, key AI Studio, dan login Google nyata ditandai **belum dilakukan**; pengujian memakai Auth Emulator dan mock; langkah nyata didokumentasikan.'],
    ['2', 'Situs dokumentasi resmi (ai.google.dev, firebase.google.com) tidak dapat dibuka langsung.',
     'Versi paket diperiksa dari registry npm/pub.dev; informasi SDK/model dari hasil pencarian. **Nama model Gemini belum diverifikasi** → dibuat konfigurabel (`GEMINI_MODEL`).'],
    ['3', '`firebase-admin` 14 memakai API modular; `admin.apps` tidak ada (TypeError saat tes pertama).',
     'Inisialisasi diganti ke `firebase-admin/app` (`initializeApp`, `getApps`) dan `firebase-admin/auth`.'],
    ['4', '`csurf` sudah diarsipkan (tidak dipelihara).', 'Dibuat perlindungan CSRF sendiri (token acak di cookie HttpOnly + meta + header, perbandingan timing-safe) dan diuji.'],
    ['5', 'Flutter: `RenderFlex overflowed by 100 pixels` pada baris "Belum punya akun?" (rawan di layar sempit).', 'Baris diganti `Wrap`; tes widget diulang hingga lulus.'],
    ['6', 'Skrip E2E menyuntikkan JS gagal karena CSP aplikasi memblokir skrip inline.', 'CSP tetap ketat; hanya konteks uji tertentu memakai `bypassCSP` di Playwright.'],
    ['7', 'Popup Google tidak bisa dibuka: SDK Firebase memuat `apis.google.com/js/api.js` (diblokir).', 'Dicatat sebagai **Belum diuji**; untuk dashboard Google dipakai kredensial IdP tiruan emulator dan diberi label jelas.'],
    ['8', 'Flutter Web tidak tampil di headless Chromium: SDK JS Firebase dimuat dari `gstatic.com` (diblokir) dan locale kosong menimbulkan "Incorrect locale information provided".',
     'Skrip E2E menyajikan SDK `firebase@12.19.0` dari paket npm lewat `page.route` (kode aplikasi tidak diubah) dan memakai `locale: id-ID`.'],
    ['9', 'Mode emulator: saat halaman di-refresh, SDK memulihkan sesi dengan menembak endpoint produksi **sebelum** `useAuthEmulator` tersambung, sehingga sesi hilang.',
     'Hanya memengaruhi mode emulator. Pada pengujian, shim SDK menyambungkan emulator saat instance Auth dibuat; uji refresh lulus (F-08). Perilaku di Firebase nyata belum diuji.'],
    ['10', 'Pengujian Flutter Web (CanvasKit) sulit: teks ada di `aria-label` pohon semantik dan `fill()` tidak andal pada field berisi.',
     'Pembaca teks memeriksa `aria-label`/`textContent`; pengisian field memverifikasi nilai `<input>` aktif dan mengulang bila tidak cocok.'],
], [0.9, 7.6, 7.5])

# ---------- 13. KESIMPULAN ----------
doc.add_heading('13. Kesimpulan dan Keterbatasan', 1)
para('Dua website (Node.js dan Flutter Web) dengan register, login email/password, login Google, dashboard terproteksi, logout, '
     'dan fitur Tanya Gemini telah dibangun. Autentikasi dikelola Firebase; backend Node.js memverifikasi identitas (session cookie untuk '
     'website Node.js, Bearer ID token untuk Flutter) sebelum memanggil Gemini, dan API key hanya berada di server. '
     'Seluruh tes otomatis dan skenario E2E yang dapat dijalankan di lingkungan ini lulus.')
para('**Keterbatasan yang masih ada:**', after=2)
bullets([
    'Belum ada proyek Firebase nyata, key AI Studio nyata, jawaban Gemini nyata, dan login Google nyata (lihat Bab 11.5 untuk daftar bukti yang kurang).',
    'Nama model Gemini default (`gemini-3.5-flash`) belum diverifikasi terhadap akun; sesuaikan `GEMINI_MODEL`.',
    'Logout hanya menghapus cookie di browser; sesi tidak dicabut di server (`revokeRefreshTokens` sengaja tidak dipakai karena akan ikut mengeluarkan sesi Flutter pengguna yang sama). '
    'Cookie yang tercuri tetap valid sampai kedaluwarsa (maks. `SESSION_MAX_DAYS`).',
    'Jika backend berada di balik reverse proxy HTTPS, perlu `app.set("trust proxy", ...)` dan `NODE_ENV=production` agar flag `Secure` bekerja.',
    'Verifikasi email tidak diwajibkan; belum ada fitur reset password.',
    'Rate limit disimpan di memori proses (tidak terbagi antar-instans).',
    'Dokumen ini berformat .docx, bukan Google Docs asli.',
])

# ---------- 14. REFERENSI ----------
doc.add_heading('14. Referensi Dokumentasi Resmi', 1)
para('Catatan: situs di bawah tidak dapat dibuka langsung dari lingkungan pengembangan; tautan dicantumkan sebagai rujukan yang perlu Anda periksa kembali.', italic=True, size=9.5)
bullets([
    'Firebase — Manage session cookies: https://firebase.google.com/docs/auth/admin/manage-cookies',
    'Firebase — Flutter setup: https://firebase.google.com/docs/flutter/setup',
    'Firebase — Autentikasi federasi (Google) di Flutter: https://firebase.google.com/docs/auth/flutter/federated-auth',
    'Firebase — Auth Emulator: https://firebase.google.com/docs/emulator-suite/connect_auth',
    'Gemini API — API keys: https://ai.google.dev/gemini-api/docs/api-key',
    'Gemini API — Getting started / quickstart: https://ai.google.dev/gemini-api/docs/get-started',
    'Paket: https://www.npmjs.com/package/@google/genai · https://www.npmjs.com/package/firebase-admin · https://pub.dev/packages/firebase_auth',
])

# ---------- footer nomor halaman ----------
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run('Halaman ')
r.font.size = Pt(8.5)
run = fp.add_run()
run.font.size = Pt(8.5)
for t_, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
    if t_:
        el = OxmlElement('w:fldChar')
        el.set(qn('w:fldCharType'), t_)
    else:
        el = OxmlElement('w:instrText')
        el.set(qn('xml:space'), 'preserve')
        el.text = txt
    run._r.append(el)

doc.core_properties.title = TITLE
doc.core_properties.author = ''
doc.save(OUT)
print('OK', OUT, 'gambar:', fig_no[0])
